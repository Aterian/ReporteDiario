# [MOD-04] liquidacion_export.py - Generación de Informe de Liquidación en Excel con openpyxl
import os
from datetime import datetime
from typing import Any
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# Paleta corporativa de Ingeap
COLOR_HEADER_BG = "C81E2B"      # Rojo institucional Ingeap
COLOR_HEADER_TEXT = "FFFFFF"    # Blanco
COLOR_SUBHEADER_BG = "F1F5F9"   # Slate 100 suave
COLOR_BORDER = "CBD5E1"         # Slate border
COLOR_ROW_ALT = "F8FAFC"        # Slate 50 alternante
COLOR_TOTAL_BG = "FEF2F2"       # Rojo tenue para total
COLOR_TOTAL_TEXT = "991B1B"     # Rojo oscuro

def _aplicar_borde(cell: Any, top: bool = True, bottom: bool = True, left: bool = True, right: bool = True):
    thin = Side(border_style="thin", color=COLOR_BORDER)
    cell.border = Border(
        top=thin if top else None,
        bottom=thin if bottom else None,
        left=thin if left else None,
        right=thin if right else None
    )

def generar_excel_informe_liquidacion(datos: dict, ruta_archivo: str) -> dict:
    """
    Genera un archivo Excel (.xlsx) corporativo con el informe de liquidación
    del empleado auditado por RRHH. Incluye:
      1. Hoja 'Resumen Liquidación': Datos del empleado, categorías de liquidación,
         días asignados, tarifas, subtotales, total a liquidar y resumen de proyectos imputados.
      2. Hoja 'Detalle Diario': Desglose de cada jornada del mes auditado.
    """
    try:
        wb = openpyxl.Workbook()
        ws_resumen = wb.active
        ws_resumen.title = "Resumen Liquidación"
        ws_resumen.views.sheetView[0].showGridLines = True

        empleado = str(datos.get("empleado") or "Colaborador").strip()
        dni = str(datos.get("dni") or "").strip()
        email = str(datos.get("email") or "").strip()
        area = str(datos.get("area") or "").strip()
        nombre_mes = str(datos.get("nombre_mes") or "Mes").strip()
        anio = str(datos.get("anio") or datetime.now().year)
        fecha_emision = str(datos.get("fecha_emision") or datetime.now().strftime("%d/%m/%Y %H:%M"))
        total_liquidar = float(datos.get("total_liquidar") or 0.0)

        # -------------------------------------------------------------
        # HOJA 1: RESUMEN DE LIQUIDACIÓN
        # -------------------------------------------------------------
        # Título principal
        ws_resumen.merge_cells("A1:E1")
        cell_titulo = ws_resumen["A1"]
        cell_titulo.value = "INGEAP S.A. • INFORME DE LIQUIDACIÓN DE SERVICIOS"
        cell_titulo.font = Font(name="Calibri", size=14, bold=True, color=COLOR_HEADER_TEXT)
        cell_titulo.fill = PatternFill(start_color=COLOR_HEADER_BG, end_color=COLOR_HEADER_BG, fill_type="solid")
        cell_titulo.alignment = Alignment(horizontal="center", vertical="center")
        ws_resumen.row_dimensions[1].height = 32

        # Subtítulo período
        ws_resumen.merge_cells("A2:E2")
        cell_periodo = ws_resumen["A2"]
        cell_periodo.value = f"Período: {nombre_mes} {anio}  |  Fecha de Emisión: {fecha_emision}"
        cell_periodo.font = Font(name="Calibri", size=10, italic=True, color="475569")
        cell_periodo.alignment = Alignment(horizontal="center", vertical="center")
        ws_resumen.row_dimensions[2].height = 20

        # Bloque de datos del colaborador
        ws_resumen.cell(row=4, column=1, value="Colaborador:").font = Font(name="Calibri", bold=True, color="1E293B")
        ws_resumen.cell(row=4, column=2, value=empleado).font = Font(name="Calibri", bold=True, size=11, color="0F172A")
        ws_resumen.cell(row=4, column=4, value="DNI:").font = Font(name="Calibri", bold=True, color="1E293B")
        ws_resumen.cell(row=4, column=5, value=dni or "-").font = Font(name="Calibri", color="0F172A")

        ws_resumen.cell(row=5, column=1, value="Área:").font = Font(name="Calibri", bold=True, color="1E293B")
        ws_resumen.cell(row=5, column=2, value=area or "-").font = Font(name="Calibri", color="0F172A")
        ws_resumen.cell(row=5, column=4, value="Email:").font = Font(name="Calibri", bold=True, color="1E293B")
        ws_resumen.cell(row=5, column=5, value=email or "-").font = Font(name="Calibri", color="0F172A")

        # Tabla de Conceptos de Liquidación
        row_idx = 7
        headers_tabla = ["Concepto / Categoría", "Días Computados", "Tarifa Diaria ($)", "Subtotal ($)"]
        ws_resumen.row_dimensions[row_idx].height = 24

        for col_idx, h_text in enumerate(headers_tabla, start=1):
            c = ws_resumen.cell(row=row_idx, column=col_idx, value=h_text)
            c.font = Font(name="Calibri", size=11, bold=True, color=COLOR_HEADER_TEXT)
            c.fill = PatternFill(start_color=COLOR_HEADER_BG, end_color=COLOR_HEADER_BG, fill_type="solid")
            c.alignment = Alignment(horizontal="center" if col_idx > 1 else "left", vertical="center")
            _aplicar_borde(c)

        categorias = datos.get("resumen_categorias") or []
        row_idx += 1
        total_dias_calc = 0.0

        for cat in categorias:
            ws_resumen.row_dimensions[row_idx].height = 20
            dias_val = float(cat.get("dias") or 0.0)
            tarifa_val = float(cat.get("tarifa_diaria") or 0.0)
            subtotal_val = float(cat.get("subtotal") or (dias_val * tarifa_val))
            total_dias_calc += dias_val

            c1 = ws_resumen.cell(row=row_idx, column=1, value=cat.get("concepto", ""))
            c2 = ws_resumen.cell(row=row_idx, column=2, value=dias_val)
            c3 = ws_resumen.cell(row=row_idx, column=3, value=tarifa_val)
            c4 = ws_resumen.cell(row=row_idx, column=4, value=subtotal_val)

            c1.font = Font(name="Calibri", size=10)
            c2.font = Font(name="Calibri", size=10, bold=True)
            c3.font = Font(name="Calibri", size=10)
            c4.font = Font(name="Calibri", size=10, bold=True)

            c1.alignment = Alignment(horizontal="left", vertical="center")
            c2.alignment = Alignment(horizontal="center", vertical="center")
            c3.alignment = Alignment(horizontal="right", vertical="center")
            c4.alignment = Alignment(horizontal="right", vertical="center")

            c2.number_format = "0.0"
            c3.number_format = "$ #,##0.00"
            c4.number_format = "$ #,##0.00"

            for cell_item in [c1, c2, c3, c4]:
                _aplicar_borde(cell_item)

            row_idx += 1

        # Fila TOTAL A LIQUIDAR
        ws_resumen.row_dimensions[row_idx].height = 26
        c_tot_label = ws_resumen.cell(row=row_idx, column=1, value="TOTAL A LIQUIDAR")
        c_tot_dias = ws_resumen.cell(row=row_idx, column=2, value=round(total_dias_calc, 1))
        c_tot_empty = ws_resumen.cell(row=row_idx, column=3, value="")
        c_tot_monto = ws_resumen.cell(row=row_idx, column=4, value=total_liquidar)

        for c_t in [c_tot_label, c_tot_dias, c_tot_empty, c_tot_monto]:
            c_t.font = Font(name="Calibri", size=11, bold=True, color=COLOR_TOTAL_TEXT)
            c_t.fill = PatternFill(start_color=COLOR_TOTAL_BG, end_color=COLOR_TOTAL_BG, fill_type="solid")
            _aplicar_borde(c_t)

        c_tot_label.alignment = Alignment(horizontal="left", vertical="center")
        c_tot_dias.alignment = Alignment(horizontal="center", vertical="center")
        c_tot_dias.number_format = "0.0"
        c_tot_monto.alignment = Alignment(horizontal="right", vertical="center")
        c_tot_monto.number_format = "$ #,##0.00"

        # Tabla de Proyectos Imputados en el Mes
        proyectos = datos.get("proyectos_imputados") or []
        if proyectos:
            row_idx += 3
            ws_resumen.cell(row=row_idx, column=1, value="PROYECTOS Y TAREAS IMPUTADAS EN EL MES").font = Font(name="Calibri", size=11, bold=True, color="0F172A")
            row_idx += 1

            headers_proy = ["Proyecto / Servicio", "Días Registrados", "Horas Totales", "Monto Imputado ($)"]
            ws_resumen.row_dimensions[row_idx].height = 22
            for col_idx, h_text in enumerate(headers_proy, start=1):
                c = ws_resumen.cell(row=row_idx, column=col_idx, value=h_text)
                c.font = Font(name="Calibri", size=10, bold=True, color="334155")
                c.fill = PatternFill(start_color=COLOR_SUBHEADER_BG, end_color=COLOR_SUBHEADER_BG, fill_type="solid")
                c.alignment = Alignment(horizontal="center" if col_idx in [2, 3] else ("right" if col_idx == 4 else "left"), vertical="center")
                _aplicar_borde(c)

            row_idx += 1
            for p in proyectos:
                ws_resumen.row_dimensions[row_idx].height = 19
                cp1 = ws_resumen.cell(row=row_idx, column=1, value=p.get("proyecto", ""))
                cp2 = ws_resumen.cell(row=row_idx, column=2, value=float(p.get("dias") or 0.0))
                cp3 = ws_resumen.cell(row=row_idx, column=3, value=float(p.get("horas") or 0.0))
                cp4 = ws_resumen.cell(row=row_idx, column=4, value=float(p.get("monto") or 0.0))

                cp1.font = Font(name="Calibri", size=9)
                cp2.font = Font(name="Calibri", size=9)
                cp3.font = Font(name="Calibri", size=9, bold=True)
                cp4.font = Font(name="Calibri", size=9)

                cp1.alignment = Alignment(horizontal="left", vertical="center")
                cp2.alignment = Alignment(horizontal="center", vertical="center")
                cp3.alignment = Alignment(horizontal="center", vertical="center")
                cp4.alignment = Alignment(horizontal="right", vertical="center")

                cp2.number_format = "0.0"
                cp3.number_format = "0.0 hs"
                cp4.number_format = "$ #,##0.00"

                for cp in [cp1, cp2, cp3, cp4]:
                    _aplicar_borde(cp)
                row_idx += 1

        # -------------------------------------------------------------
        # HOJA 2: DETALLE DIARIO CRONOLÓGICO
        # -------------------------------------------------------------
        ws_detalle = wb.create_sheet(title="Detalle Diario")
        ws_detalle.views.sheetView[0].showGridLines = True

        ws_detalle.merge_cells("A1:G1")
        c_det_top = ws_detalle["A1"]
        c_det_top.value = f"DETALLE DE REGISTROS DIARIOS • {empleado.upper()} ({nombre_mes.upper()} {anio})"
        c_det_top.font = Font(name="Calibri", size=12, bold=True, color=COLOR_HEADER_TEXT)
        c_det_top.fill = PatternFill(start_color=COLOR_HEADER_BG, end_color=COLOR_HEADER_BG, fill_type="solid")
        c_det_top.alignment = Alignment(horizontal="center", vertical="center")
        ws_detalle.row_dimensions[1].height = 28

        headers_det = ["Fecha", "Día", "Modalidad (Lugar)", "Proyecto / Servicio", "Horas", "Tipo Costo", "Importe ($)"]
        ws_detalle.row_dimensions[2].height = 22
        for col_idx, h_text in enumerate(headers_det, start=1):
            c = ws_detalle.cell(row=2, column=col_idx, value=h_text)
            c.font = Font(name="Calibri", size=10, bold=True, color="334155")
            c.fill = PatternFill(start_color=COLOR_SUBHEADER_BG, end_color=COLOR_SUBHEADER_BG, fill_type="solid")
            c.alignment = Alignment(horizontal="center" if col_idx in [1, 2, 5, 6] else ("right" if col_idx == 7 else "left"), vertical="center")
            _aplicar_borde(c)

        registros = datos.get("registros_detalle") or []
        r_det_idx = 3
        for r in registros:
            ws_detalle.row_dimensions[r_det_idx].height = 19
            fec = str(r.get("fecha") or "")
            dia_s = str(r.get("dia_semana") or "")
            mod = str(r.get("tipo_ocf") or r.get("lugar") or "")
            srv = str(r.get("servicio") or "")
            hrs = float(r.get("horas") or 0.0)
            t_costo = str(r.get("tipo_costo") or "")
            c_dia = float(r.get("costo_dia") or 0.0)

            cd1 = ws_detalle.cell(row=r_det_idx, column=1, value=fec)
            cd2 = ws_detalle.cell(row=r_det_idx, column=2, value=dia_s)
            cd3 = ws_detalle.cell(row=r_det_idx, column=3, value=mod)
            cd4 = ws_detalle.cell(row=r_det_idx, column=4, value=srv)
            cd5 = ws_detalle.cell(row=r_det_idx, column=5, value=hrs)
            cd6 = ws_detalle.cell(row=r_det_idx, column=6, value=t_costo)
            cd7 = ws_detalle.cell(row=r_det_idx, column=7, value=c_dia)

            for cell_d in [cd1, cd2, cd3, cd4, cd5, cd6, cd7]:
                cell_d.font = Font(name="Calibri", size=9)
                _aplicar_borde(cell_d)

            cd1.alignment = Alignment(horizontal="center", vertical="center")
            cd2.alignment = Alignment(horizontal="center", vertical="center")
            cd3.alignment = Alignment(horizontal="left", vertical="center")
            cd4.alignment = Alignment(horizontal="left", vertical="center")
            cd5.alignment = Alignment(horizontal="center", vertical="center")
            cd6.alignment = Alignment(horizontal="center", vertical="center")
            cd7.alignment = Alignment(horizontal="right", vertical="center")

            cd5.number_format = "0.0 hs"
            cd7.number_format = "$ #,##0.00"

            r_det_idx += 1

        # Autoajuste de anchos de columnas en ambas hojas
        for ws in [ws_resumen, ws_detalle]:
            for col in ws.columns:
                max_len = 0
                col_letter = get_column_letter(col[0].column)
                for cell in col:
                    if cell.value is not None:
                        # Si es celda combinada de título, ignorar longitud para no sobredimensionar la primera columna
                        if cell.row in [1, 2]:
                            continue
                        val_str = str(cell.value)
                        if len(val_str) > max_len:
                            max_len = len(val_str)
                ws.column_dimensions[col_letter].width = max(max_len + 4, 14)

        ws_resumen.column_dimensions["A"].width = 38
        ws_resumen.column_dimensions["B"].width = 22
        ws_resumen.column_dimensions["C"].width = 20
        ws_resumen.column_dimensions["D"].width = 24
        ws_resumen.column_dimensions["E"].width = 24

        ws_detalle.column_dimensions["A"].width = 14
        ws_detalle.column_dimensions["B"].width = 12
        ws_detalle.column_dimensions["C"].width = 22
        ws_detalle.column_dimensions["D"].width = 44
        ws_detalle.column_dimensions["E"].width = 12
        ws_detalle.column_dimensions["F"].width = 16
        ws_detalle.column_dimensions["G"].width = 18

        wb.save(ruta_archivo)
        return {
            "exito": True,
            "ruta": ruta_archivo,
            "mensaje": f"Informe de liquidación guardado exitosamente en:\n{ruta_archivo}"
        }
    except Exception as e:
        print(f"[LiquidacionExport] Error al generar Excel: {e}")
        return {"exito": False, "error": str(e)}
