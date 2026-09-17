import os
import calendar
from datetime import datetime, date, timedelta
import openpyxl
from typing import Any
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from database import obtener_rosters

# Paleta corporativa de Ingeap y estilos
COLOR_HEADER_BG = "C81E2B"      # Rojo institucional Ingeap
COLOR_HEADER_TEXT = "FFFFFF"    # Blanco
COLOR_SUNDAY_BG = "FEF3C7"      # Fondo suave dorado/ámbar para domingos
COLOR_SUNDAY_HEADER = "F59E0B"  # Ámbar destacado para cabecera domingo
COLOR_CAMPO_CELL = "D1FAE5"     # Verde suave para Campo
COLOR_CAMPO_TEXT = "065F46"     # Verde oscuro
COLOR_FRANCO_CELL = "EDE9FE"    # Púrpura suave para Franco
COLOR_FRANCO_TEXT = "5B21B6"    # Púrpura oscuro
COLOR_BORDER = "CBD5E1"         # Slate border

def _aplicar_borde(cell: Any, top: bool = True, bottom: bool = True, left: bool = True, right: bool = True):
    thin = Side(border_style="thin", color=COLOR_BORDER)
    cell.border = Border(
        top=thin if top else None,
        bottom=thin if bottom else None,
        left=thin if left else None,
        right=thin if right else None
    )

def generar_excel_roster_mes(anio: int, mes: int, ruta_archivo: str, proyecto: str | None = None) -> dict:
    """
    Genera un archivo .xlsx profesional para el mes y proyecto especificado,
    dividido en 3 hojas:
      1. Ciclo Completo (días 1 al fin de mes)
      2. 1ra Quincena (días 1 al 15)
      3. 2da Quincena (días 16 al fin de mes)
    Resalta los domingos en todas las tablas y realiza los cálculos de días y totales por proyecto.
    """
    try:
        wb = openpyxl.Workbook()
        # Eliminar la hoja por defecto
        wb.remove(wb.active)

        _, dias_en_mes = calendar.monthrange(anio, mes)
        nombres_meses = [
            "", "Enero", "Febrero", "Marzo", "Abril", "Mayo", "Junio",
            "Julio", "Agosto", "Septiembre", "Octubre", "Noviembre", "Diciembre"
        ]
        nombre_mes = nombres_meses[mes]

        # Fechas límite del mes
        fecha_ini_mes_str = f"{anio:04d}-{mes:02d}-01"
        fecha_fin_mes_str = f"{anio:04d}-{mes:02d}-{dias_en_mes:02d}"

        # Obtener todos los registros de roster que toquen este mes
        rosters_raw = obtener_rosters(fecha_ini_mes_str, fecha_fin_mes_str)

        # Filtrar exclusivamente por el proyecto indicado si fue provisto
        proy_filtro = proyecto.strip() if (proyecto and isinstance(proyecto, str)) else ""
        if proy_filtro and proy_filtro.upper() != "TODOS":
            rosters_raw = [r for r in rosters_raw if r.get("proyecto", "").strip() == proy_filtro]

        # Definición de las 3 hojas
        secciones = [
            ("Ciclo Completo", 1, dias_en_mes),
            ("1ra Quincena", 1, min(15, dias_en_mes)),
            ("2da Quincena", 16, dias_en_mes)
        ]

        dias_semana_abrev = ["Lun", "Mar", "Mié", "Jue", "Vie", "Sáb", "Dom"]

        # Agrupar empleados y proyectos únicos
        empleados_proyectos = set()
        for r in rosters_raw:
            empleados_proyectos.add((r["empleado"].strip(), r["proyecto"].strip()))
        
        lista_emp_proy = sorted(list(empleados_proyectos), key=lambda x: (x[0], x[1]))

        for titulo_hoja, dia_ini, dia_fin in secciones:
            ws: Any = wb.create_sheet(title=titulo_hoja)
            ws.views.sheetView[0].showGridLines = True

            # Título principal de la hoja
            ws.merge_cells("A1:I1")
            celda_titulo = ws["A1"]
            celda_titulo.value = f"REPORTE DE ROSTER - {titulo_hoja.upper()} ({nombre_mes.upper()} {anio})"
            celda_titulo.font = Font(name="Calibri", size=14, bold=True, color="FFFFFF")
            celda_titulo.fill = PatternFill(start_color=COLOR_HEADER_BG, end_color=COLOR_HEADER_BG, fill_type="solid")
            celda_titulo.alignment = Alignment(horizontal="center", vertical="center")
            ws.row_dimensions[1].height = 32

            # Subtítulo informativo con indicación del proyecto
            ws.merge_cells("A2:I2")
            celda_sub = ws["A2"]
            if proy_filtro and proy_filtro.upper() != "TODOS":
                celda_sub.value = f"Proyecto: {proy_filtro} | Período: {dia_ini:02d}/{mes:02d}/{anio} al {dia_fin:02d}/{mes:02d}/{anio} | Ingeap - Recursos Humanos"
            else:
                celda_sub.value = f"Período: {dia_ini:02d}/{mes:02d}/{anio} al {dia_fin:02d}/{mes:02d}/{anio} | Ingeap - Recursos Humanos"
            celda_sub.font = Font(name="Calibri", size=10, italic=True, color="475569")
            celda_sub.alignment = Alignment(horizontal="left", vertical="center")
            ws.row_dimensions[2].height = 20

            # Encabezados de columnas fijas (Fila 4)
            headers_fijos = [
                ("A4", "Empleado", 25),
                ("B4", "Proyecto Asignado", 32),
                ("C4", "Días Norm. Activos", 16),
                ("D4", "Domingos Activos", 16),
                ("E4", "Precio Día Norm. ($)", 18),
                ("F4", "Precio Domingo ($)", 18),
                ("G4", "Total Normal ($)", 18),
                ("H4", "Total Domingos ($)", 18),
                ("I4", "Total Liquidar ($)", 19),
            ]

            fill_header = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid") # Slate oscuro
            font_header = Font(name="Calibri", size=10, bold=True, color="FFFFFF")

            for pos, texto, ancho in headers_fijos:
                c = ws[pos]
                c.value = texto
                c.font = font_header
                c.fill = fill_header
                c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                col_letra = pos[0]
                ws.column_dimensions[col_letra].width = ancho
                _aplicar_borde(c)

            # Encabezados de días del calendario (a partir de la columna J)
            col_idx = 10
            for d in range(dia_ini, dia_fin + 1):
                fecha_dia = date(anio, mes, d)
                es_domingo = fecha_dia.weekday() == 6
                nombre_dia = dias_semana_abrev[fecha_dia.weekday()]

                col_letra = get_column_letter(col_idx)
                c_dia = ws[f"{col_letra}4"]
                c_dia.value = f"{nombre_dia}\n{d}"
                c_dia.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                
                if es_domingo:
                    c_dia.fill = PatternFill(start_color=COLOR_SUNDAY_HEADER, end_color=COLOR_SUNDAY_HEADER, fill_type="solid")
                    c_dia.font = Font(name="Calibri", size=9, bold=True, color="FFFFFF")
                else:
                    c_dia.fill = PatternFill(start_color="334155", end_color="334155", fill_type="solid")
                    c_dia.font = Font(name="Calibri", size=9, bold=True, color="FFFFFF")
                
                ws.column_dimensions[col_letra].width = 6.5
                _aplicar_borde(c_dia)
                col_idx += 1

            ws.row_dimensions[4].height = 28

            # Filas de datos
            fila_actual = 5
            total_general_quincena = 0.0

            if not lista_emp_proy:
                # Fila sin datos
                ws.merge_cells(start_row=5, start_column=1, end_row=5, end_column=col_idx - 1)
                c_vacio = ws.cell(row=5, column=1)
                if proy_filtro and proy_filtro.upper() != "TODOS":
                    c_vacio.value = f"No hay registros de roster cargados para el proyecto '{proy_filtro}' en este período."
                else:
                    c_vacio.value = "No hay registros de roster cargados para este período."
                c_vacio.alignment = Alignment(horizontal="center", vertical="center")
                c_vacio.font = Font(name="Calibri", size=11, italic=True, color="64748B")
                ws.row_dimensions[5].height = 24
                fila_actual = 6
            else:
                for emp_nom, proy_nom in lista_emp_proy:
                    dias_norm_activos = 0
                    domingos_activos = 0
                    precio_dia_usado = 0.0
                    precio_dom_usado = 0.0

                    # Matriz de días para esta persona y proyecto
                    estado_por_dia = {} # d -> ('Campo' o 'Franco', precio_dia, precio_dom)

                    for d in range(dia_ini, dia_fin + 1):
                        fecha_str = f"{anio:04d}-{mes:02d}-{d:02d}"
                        
                        # Buscar si coincide con algún registro
                        for r in rosters_raw:
                            if r["empleado"].strip() == emp_nom and r["proyecto"].strip() == proy_nom:
                                if r["fecha_inicio"] <= fecha_str <= r["fecha_fin"]:
                                    estado_por_dia[d] = (r["tipo"], float(r.get("precio_dia") or 0), float(r.get("precio_domingo") or 0))
                                    break

                    # Calcular conteos y tarifas
                    for d, (tipo, pdia, pdom) in estado_por_dia.items():
                        fecha_d = date(anio, mes, d)
                        es_domingo = fecha_d.weekday() == 6
                        if tipo == "Campo":
                            if es_domingo:
                                domingos_activos += 1
                                if pdom > 0: precio_dom_usado = pdom
                            else:
                                dias_norm_activos += 1
                                if pdia > 0: precio_dia_usado = pdia

                    # Si no se usó tarifa en el cálculo pero hay un roster registrado para este proyecto, tomar la última cargada
                    if precio_dia_usado == 0 or precio_dom_usado == 0:
                        for r in rosters_raw:
                            if r["empleado"].strip() == emp_nom and r["proyecto"].strip() == proy_nom:
                                if precio_dia_usado == 0 and float(r.get("precio_dia") or 0) > 0:
                                    precio_dia_usado = float(r["precio_dia"])
                                if precio_dom_usado == 0 and float(r.get("precio_domingo") or 0) > 0:
                                    precio_dom_usado = float(r["precio_domingo"])

                    tot_normal = dias_norm_activos * precio_dia_usado
                    tot_domingos = domingos_activos * precio_dom_usado
                    tot_liquidar = tot_normal + tot_domingos
                    total_general_quincena += tot_liquidar

                    # Escribir columnas fijas
                    ws.cell(row=fila_actual, column=1, value=emp_nom).alignment = Alignment(horizontal="left", vertical="center")
                    ws.cell(row=fila_actual, column=2, value=proy_nom).alignment = Alignment(horizontal="left", vertical="center")
                    
                    c_dna = ws.cell(row=fila_actual, column=3, value=dias_norm_activos)
                    c_dna.alignment = Alignment(horizontal="center", vertical="center")
                    
                    c_doma = ws.cell(row=fila_actual, column=4, value=domingos_activos)
                    c_doma.alignment = Alignment(horizontal="center", vertical="center")

                    c_pd = ws.cell(row=fila_actual, column=5, value=precio_dia_usado)
                    c_pd.number_format = '$#,##0.00'
                    c_pd.alignment = Alignment(horizontal="right", vertical="center")

                    c_pdom = ws.cell(row=fila_actual, column=6, value=precio_dom_usado)
                    c_pdom.number_format = '$#,##0.00'
                    c_pdom.alignment = Alignment(horizontal="right", vertical="center")

                    c_tn = ws.cell(row=fila_actual, column=7, value=tot_normal)
                    c_tn.number_format = '$#,##0.00'
                    c_tn.alignment = Alignment(horizontal="right", vertical="center")

                    c_tdom = ws.cell(row=fila_actual, column=8, value=tot_domingos)
                    c_tdom.number_format = '$#,##0.00'
                    c_tdom.alignment = Alignment(horizontal="right", vertical="center")

                    c_tliq = ws.cell(row=fila_actual, column=9, value=tot_liquidar)
                    c_tliq.number_format = '$#,##0.00'
                    c_tliq.font = Font(name="Calibri", size=10, bold=True, color="1E293B")
                    c_tliq.alignment = Alignment(horizontal="right", vertical="center")

                    for col_f in range(1, 10):
                        _aplicar_borde(ws.cell(row=fila_actual, column=col_f))

                    # Escribir celdas de días
                    c_idx = 10
                    for d in range(dia_ini, dia_fin + 1):
                        fecha_dia = date(anio, mes, d)
                        es_domingo = fecha_dia.weekday() == 6
                        c_dia_val = ws.cell(row=fila_actual, column=c_idx)
                        
                        info_dia = estado_por_dia.get(d)
                        if info_dia:
                            tipo = info_dia[0]
                            if tipo == "Campo":
                                c_dia_val.value = "C"
                                c_dia_val.fill = PatternFill(start_color=COLOR_CAMPO_CELL, end_color=COLOR_CAMPO_CELL, fill_type="solid")
                                c_dia_val.font = Font(name="Calibri", size=9, bold=True, color=COLOR_CAMPO_TEXT)
                            else:
                                c_dia_val.value = "F"
                                c_dia_val.fill = PatternFill(start_color=COLOR_FRANCO_CELL, end_color=COLOR_FRANCO_CELL, fill_type="solid")
                                c_dia_val.font = Font(name="Calibri", size=9, bold=True, color=COLOR_FRANCO_TEXT)
                        else:
                            c_dia_val.value = ""
                            if es_domingo:
                                c_dia_val.fill = PatternFill(start_color=COLOR_SUNDAY_BG, end_color=COLOR_SUNDAY_BG, fill_type="solid")

                        c_dia_val.alignment = Alignment(horizontal="center", vertical="center")
                        _aplicar_borde(c_dia_val)
                        c_idx += 1

                    ws.row_dimensions[fila_actual].height = 20
                    fila_actual += 1

                # Fila de Totales Generales
                ws.merge_cells(start_row=fila_actual, start_column=1, end_row=fila_actual, end_column=8)
                c_tot_label = ws.cell(row=fila_actual, column=1)
                c_tot_label.value = "TOTAL GENERAL A LIQUIDAR:"
                c_tot_label.font = Font(name="Calibri", size=10, bold=True, color="1E293B")
                c_tot_label.alignment = Alignment(horizontal="right", vertical="center")
                
                c_tot_val = ws.cell(row=fila_actual, column=9)
                c_tot_val.value = total_general_quincena
                c_tot_val.font = Font(name="Calibri", size=11, bold=True, color="C81E2B")
                c_tot_val.number_format = '$#,##0.00'
                c_tot_val.alignment = Alignment(horizontal="right", vertical="center")
                c_tot_val.fill = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")

                for c_pos in range(1, col_idx):
                    _aplicar_borde(ws.cell(row=fila_actual, column=c_pos))

                ws.row_dimensions[fila_actual].height = 24
                fila_actual += 2

                # Leyenda al pie
                c_leyenda = ws.cell(row=fila_actual, column=1)
                c_leyenda.value = "Referencias: [ C = Día en Campo/Obra ]   [ F = Día de Franco ]   [ Columnas Ámbar = Domingos ]"
                c_leyenda.font = Font(name="Calibri", size=9, italic=True, color="64748B")

        wb.save(ruta_archivo)
        return {"exito": True, "ruta": ruta_archivo, "mensaje": f"Archivo exportado correctamente en: {ruta_archivo}"}
    except Exception as e:
        print(f"[RosterExport] Error exportando Excel: {e}")
        return {"exito": False, "error": str(e)}
